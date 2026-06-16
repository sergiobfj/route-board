from fastapi import APIRouter, Depends, File, UploadFile
from sqlmodel import Session, select
from models import Route, RouteStop, RouteItem, RouteCreate
from database import get_session

import openpyxl
from io import BytesIO

router = APIRouter()

@router.post("/routes")
def route_create(route: RouteCreate, session: Session = Depends(get_session)):

    db_route = Route(name=route.name)
    session.add(db_route)
    session.commit()
    session.refresh(db_route)

    for stop in route.stops:
        db_stop = RouteStop(city=stop.city, route_id=db_route.id)
        session.add(db_stop)
        session.commit()
        session.refresh(db_stop)

        for item in stop.items:
            db_item = RouteItem(
                module_qty=item.module_qty,
                module_brand=item.module_brand,
                inverter_qty=item.inverter_qty,
                inverter_brand=item.inverter_brand,
                stop_id=db_stop.id
            )
            session.add(db_item)

    session.commit()
    return db_route

@router.get("/routes")
def list_routes(session: Session = Depends(get_session)):

    routes = session.exec(select(Route)).all()

    result = []

    for route in routes:
        stops = session.exec(select(RouteStop).where(RouteStop.route_id == route.id)).all()

        stops_with_items = []
        for stop in stops:
            items = session.exec(select(RouteItem).where(RouteItem.stop_id == stop.id)).all()
            stops_with_items.append ({
                "city": stop.city,
                "items": items
            })

        result.append({
            "id": route.id,
            "name": route.name,
            "status": route.status,
            "driver_name": route.driver_name,
            "truck_plate": route.truck_plate,
            "stops": stops_with_items
        })

    return result

@router.patch("/routes/{id}")
def update_route(id: int, status: str, session: Session = Depends(get_session)):

    route = session.get(Route, id)
    route.status = status

    session.add(route)
    session.commit()
    session.refresh(route)
    
    return route

@router.get("/display")
def get_display(session: Session = Depends(get_session)):

    statement = select(Route).where(Route.status == "in_progress").order_by(Route.created_at)
    in_progress = session.exec(statement).first()

    if in_progress:
        stops = session.exec(select(RouteStop).where(RouteStop.route_id == in_progress.id)).all()
        
        stops_with_items = []
        for stop in stops:
            items = session.exec(select(RouteItem).where(RouteItem.stop_id == stop.id)).all()
            stops_with_items.append({
                "city": stop.city,
                "items": items
            })
        
        in_progress_data = {
            "id": in_progress.id,
            "name": in_progress.name,
            "stops": stops_with_items
        }
    else:
        in_progress_data = None


    statement = select(Route).where(Route.status == "waiting").order_by(Route.created_at)
    waiting = session.exec(statement).first()

    if waiting:
        stops = session.exec(select(RouteStop).where(RouteStop.route_id == waiting.id)).all()

        stops_with_items = []
        for stop in stops:
            items = session.exec(select(RouteItem).where(RouteItem.stop_id == stop.id)).all()
            stops_with_items.append ({
                "city": stop.city,
                "items": items
            })

        waiting_data = {
            "id": waiting.id,
            "name": waiting.name,
            "driver_name": waiting.driver_name,
            "truck_plate": waiting.truck_plate,
            "stops": stops_with_items
        }

    else:
        waiting_data = None

    
    return {
        "in_progress": in_progress_data,
        "next": waiting_data
    }

            
@router.post("/import")
async def import_route(file: UploadFile, session: Session = Depends(get_session)):
    contents = await file.read()
    workbook = openpyxl.load_workbook(BytesIO(contents))
    sheet = workbook.active

    header_row = list(sheet.iter_rows(min_row=4, max_row=4, values_only=True))[0]
    route_name = header_row[1]
    driver_name = header_row[3]
    truck_plate = header_row[5]

    existing = session.exec(select(Route).where(Route.name == route_name)).first()
    if existing:
        return {"message": "Rota já existe", "route_id": existing.id}

    db_route = Route(name=route_name, driver_name=driver_name, truck_plate=truck_plate)

    db_route = Route(name=route_name, driver_name=driver_name, truck_plate=truck_plate)
    session.add(db_route)
    session.commit()
    session.refresh(db_route)

    stops_dict = {}

    for row in sheet.iter_rows(min_row=6, values_only=True):
        if row[1] == "TOTAL" or row[1] is None:
            break

        code, client_name, qty, modules, inverter, kit_tech, city, notes, roof = row[:9]

        if city not in stops_dict:
            db_stop = RouteStop(city=city, route_id=db_route.id)
            session.add(db_stop)
            session.commit()
            session.refresh(db_stop)
            stops_dict[city] = db_stop.id

        db_item = RouteItem(
            module_qty=qty,
            module_brand=modules,
            inverter_qty=1,
            inverter_brand=inverter,
            stop_id=stops_dict[city]
        )
        session.add(db_item)

    session.commit()
    return {"message": "Importado com sucesso", "route_id": db_route.id}